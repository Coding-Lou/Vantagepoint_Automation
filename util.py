import json
from pypdf import PdfReader, PdfWriter
from pathlib import Path
import os
import requests
from openpyxl.worksheet.table import Table
from datetime import datetime

def show_welcome_banner():
    banner = rf"""
         ██████╗    ██████    █████╗ 
        ██╔═══██╗  ██        ██╔══██╗
        ██║   ██║  ██        ███████║
        ██║▄▄ ██║  ██        ██╔══██║
        ╚██████╔╝  ╚██████   ██║  ██║
         ╚══▀▀═╝    ╚═════   ╚═╝  ╚═╝
    ──────────────────────────────────────────────────────────
            🛠️  QCA Accounting Team Automation Script
            🏷️  Version: {get_config(["VERSION"])}
    ──────────────────────────────────────────────────────────
    """

    print(banner)
    print("🚀 Welcome! \n")

def get_config(klist):
    try:
        with open("config.json", "r", encoding="utf-8") as f:
            config = json.load(f)
        for key in klist:
            config = config[key]
        return config
    
    except Exception as e:
        print(f"⚠️ Get config failed: ", e)
        return None

def set_config(key, value):
    try:
        with open("config.json", "r", encoding="utf-8") as f:
            config = json.load(f)
        
        config[key] = value
        
        with open("config.json", "w", encoding="utf-8") as f:
            json.dump(config, f, indent=4, ensure_ascii=False)
            f.flush()

        print(f"💾 {key} been updated to {value}. ")
        #print()

    except Exception as e:
        print(f"⚠️ {key} updated failed: ", e)
        #print()

def check_login():
    try: 
        url = "https://qcadeltek03.qcasystems.com/vantagepoint/visionservices.asmx/GetIAccessConfiguration"
        payload = {"sessionID": get_config(["TOKEN"])}
        response = requests.post(url, headers = set_headers(), json = payload)
        cookies = get_config(["COOKIES"])
        if response.status_code == 200 and "ASP.NET_SessionId" in cookies:
            data = response.json()
            print("✅ Login Success, User: " + data["d"]["UserInfo"]["EMail"])
            print()
            return True
        else: 
            return False
    except:
        print("❌ Error in function check_login()")
        return False

def merge_amazon_invoices():
    path = input("Please input the folder path: ")
    input_folder = Path(path)  
    output_file = input_folder / "combined_output.pdf"

    combined_writer = PdfWriter()

    for pdf_file in input_folder.glob("*.pdf"):
        try:
            reader = PdfReader(str(pdf_file))
            if len(reader.pages) >= 1:
                combined_writer.add_page(reader.pages[0])
            if len(reader.pages) == 2:
                combined_writer.add_page(reader.pages[1])
            if len(reader.pages) >= 3:
                combined_writer.add_page(reader.pages[2])
            print(f"✅ Done: {pdf_file.name}")
        except Exception as e:
            print(f"⚠️ Error {pdf_file.name}: {e}")

    with open(output_file, "wb") as f:
        combined_writer.write(f)

    print(f"\n🎉 Success: {output_file}")

def merge_pdfs():
    folder_path = input("📂 Please input folder path of the pdf files: ").strip()
    output_filename = "merged_output.pdf"

    if not os.path.isdir(folder_path):
        print("❌ Invalid folder path")

    pdf_writer = PdfWriter()

    pdf_files = [
        os.path.join(folder_path, f)
        for f in os.listdir(folder_path)
        if f.lower().endswith('.pdf')
    ]

    # Sort by created time
    pdf_files.sort(key=lambda f: os.path.getctime(f))

    if not pdf_files:
        print("❌ No pdf files")
        return

    for pdf_path in pdf_files:
        try:
            reader = PdfReader(pdf_path)
            for page in reader.pages:
                pdf_writer.add_page(page)
            print(f"✅ Add: {os.path.basename(pdf_path)}")
        except Exception as e:
            print(f"⚠️ Skip {pdf_path}: {e}")

    output_path = os.path.join(folder_path, output_filename)
    with open(output_path, "wb") as out_file:
        pdf_writer.write(out_file)

    print(f"\n🎉 Success the merged pdf file: {output_path}")

def set_headers():
    WWWBEARER = get_config(["WWWBEARER"])
    TOKEN = get_config(["TOKEN"])
    COOKIES = get_config(["COOKIES"])
    headers = {
        "accept": "application/json, text/javascript, */*; q=0.01",
        "Content-Type": "application/json; charset=UTF-8",
        "www-bearer": WWWBEARER,
        "Token": TOKEN,
        "Cookie": COOKIES
    }
    return headers

def assamble_projects(projects):
    searchOptions = [{"name":"Status","value":"[IS_EMPTY]","type":"dropdown","seq":1,"tableName":"PR","opp":"!=","condition":"and","searchLevel":1,"valueDescription":""}]
    for project in projects:
        searchOptions.append({"name":"selectedResultIds","value":project,"type":"wbs1","seq":2,"searchLevel":0,"valueDescription":"Monthly Maintenance"})
    
    return searchOptions

def check_folder(folderName):
    if not os.path.exists(folderName):
        os.makedirs(folderName)
        print(f"📁 Folder created: {folderName}")

def clear_folder(folderName):
    onedrivedir = get_config(["ONEDRIVEDIR"])
    workdir = get_config(["WORKDIR"])
    folder_path = os.path.join(onedrivedir, workdir, folderName)
    for filename in os.listdir(folder_path):
        file_path = os.path.join(folder_path, filename)
        if os.path.isfile(file_path):
            os.remove(file_path)

def get_vendor_email(clientID):
    headers = set_headers()
    try:
        url = "https://qcadeltek03.qcasystems.com/vantagepoint/vision/Client/"+ clientID +"/Address/"
        response = requests.get(url, headers=headers  )
        data = response.json()
        email = ""
        if (len(data) == 1):
            email = data[0]["Email"]
        else:
            for d in data:
                if d["Email"] != "" and (not d["Email"] in email): 
                    if "AP Automation" in d['Address'] or "AP Mailing" in d['Address']: 
                        email = d["Email"]+";"
                        break
                    else:
                        email += d["Email"]+";"
        return email
    except Exception as e:
        print("❌ Error in function get_vendor_email()")

def get_clientID(clientName):
    try:
        url = "https://qcadeltek03.qcasystems.com/vantagepoint/visionservices.asmx/GetLookupHash"
        headers = set_headers()
        payload = {
            "sessionID": get_config(["TOKEN"]),
            "hash":{
                "filter": clientName,
                "lookuptype": "clientvendor","vendorOnly":"N"
            },
            "page": 1,
            "pagesize": 20,
            "order": "name"
        }
        response = requests.post(url, headers=headers, json=payload)
        records = response.json()
        for record in records['d']:
            if record['IsClient'] == "Y":
                return record['Key']
    except Exception as e:
        print("❌ Error in function get_clientID() with input: " + clientName)

def save_excel(wb, records):
    try:
        ws = wb.active
        table_range = f"A1:I{records}"
        if "Table1" in ws.tables:
            del ws.tables["Table1"]
        tab = Table(displayName = "Table1", ref = table_range)
        ws.add_table(tab)
        timestamp = datetime.now().strftime("%Y%m%d")
        excelName = os.path.join( f"Job_{timestamp}.xlsx")
        wb.save(excelName)

        return excelName
    except Exception as e:
        print("❌Error in function save_excel()")

def download_with_progress(url, save_path, latest_version):
    r = requests.get(url, stream=True)
    total = int(r.headers.get('content-length', 0))

    downloaded = 0
    chunk_size = 8192

    print("\nDownloading update...\n")

    with open(save_path, "wb") as f:
        for chunk in r.iter_content(chunk_size):
            if chunk:
                f.write(chunk)
                downloaded += len(chunk)

                percent = downloaded / total * 100 if total else 0
                bar = "█" * int(percent / 2)   # 50 chars bar
                space = " " * (50 - len(bar))

                print(f"\r[{bar}{space}] {percent:6.2f}%  ({downloaded/1024/1024:.2f} MB / {total/1024/1024:.2f} MB)", end="")

    print("\nDownload complete!\n")
    set_config("VERSION", latest_version)
