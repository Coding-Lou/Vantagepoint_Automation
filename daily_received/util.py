import json
import shutil
import time
from pypdf import PdfReader, PdfWriter
from pathlib import Path
import os
import requests
from openpyxl.worksheet.table import Table
from datetime import datetime
from dateutil.relativedelta import relativedelta
import pandas as pd
import subprocess
import glob
import sys

def show_welcome_banner():
    banner = rf"""
         ██████╗    ██████    █████╗ 
        ██╔═══██╗  ██        ██╔══██╗
        ██║   ██║  ██        ███████║
        ██║▄▄ ██║  ██        ██╔══██║
        ╚██████╔╝  ╚██████   ██║  ██║
         ╚══▀▀═╝    ╚═════   ╚═╝  ╚═╝
    ──────────────────────────────────────────────────────────
            🛠️  QCA Accounting Team Automation Script
            🏷️  Version: {get_config(["VERSION"])}
    ──────────────────────────────────────────────────────────
    """

    print(banner)
    print("🚀 Welcome! \n")

def show_menu():
    MENU = """
======================================================
                   MAIN MENU
======================================================
  1) AP - Remittance
  2) AR - Statements
  3) Report Preparation - Project Status
  4) Report Preparation - Bridge Report
  5) Shipping Monitor
  6) Daily Receiving Notification
  7) Merge PDF

  0) Exit
======================================================
"""
    print(MENU)

def get_runtime_dir() -> Path:
    """Get the runtime directory. In frozen mode, returns exe directory."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent

def get_config(klist):
    """Get config value. Config file is in the main directory (same level as main.py or exe)."""
    base_dir = get_runtime_dir()
    
    # Config is in the main directory (same level as main.py or exe)
    config_path = base_dir / "config.json"
    
    # In frozen mode, if config doesn't exist, try to copy from bundled resource
    if getattr(sys, "frozen", False) and not config_path.exists():
        try:
            if hasattr(sys, '_MEIPASS'):
                # PyInstaller temporary folder
                bundled_config = Path(sys._MEIPASS) / "config.json"
                if bundled_config.exists():
                    import shutil
                    shutil.copy2(bundled_config, config_path)
        except Exception:
            pass
    try:
        with open(config_path, "r", encoding="utf-8") as f:
            config = json.load(f)
        node = config
        for key in klist:
            if key not in node: 
                node[key] = ""
            node = node[key]

        with open(config_path, "w", encoding="utf-8") as f:
            json.dump(config, f, indent=4, ensure_ascii=False)
            
        return node
    
    except Exception as e:
        print(f"⚠️ Get config failed")
        return None

def set_config(key, value):
    """Set config value. Config file is in the main directory (same level as main.py or exe)."""
    base_dir = get_runtime_dir()
    
    # Config is in the main directory (same level as main.py or exe)
    config_path = base_dir / "config.json"
    
    # In frozen mode, if config doesn't exist, try to copy from bundled resource
    if getattr(sys, "frozen", False) and not config_path.exists():
        try:
            if hasattr(sys, '_MEIPASS'):
                bundled_config = Path(sys._MEIPASS) / "config.json"
                if bundled_config.exists():
                    import shutil
                    shutil.copy2(bundled_config, config_path)
        except Exception:
            pass
    try:
        with open(config_path, "r", encoding="utf-8") as f:
            config = json.load(f)
        
        config[key] = value

        with open(config_path, "w", encoding="utf-8") as f:
            json.dump(config, f, indent=4, ensure_ascii=False)
            f.flush()

        print(f"💾 {key} been updated to {value}. ")

    except Exception as e:
        print(f"⚠️ {key} updated failed: ", e)

def check_login():
    from pathlib import Path
    import json
    import time
    DEBUG_LOG_PATH = Path(r"c:\cursor\.cursor\debug.log")
    def _agent_log(hypothesis_id: str, location: str, message: str, data: dict = None):
        payload = {
            "sessionId": "debug-session",
            "runId": "pre-fix",
            "hypothesisId": hypothesis_id,
            "location": location,
            "message": message,
            "data": data or {},
            "timestamp": int(time.time() * 1000),
        }
        try:
            DEBUG_LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
            with DEBUG_LOG_PATH.open("a", encoding="utf-8") as f:
                f.write(json.dumps(payload, ensure_ascii=False) + "\n")
        except Exception:
            pass
    #region agent log
    _agent_log("H2", "util.check_login", "check_login() entry", {})
    #endregion
    try: 
        token = get_config(["TOKEN"])
        #region agent log
        _agent_log("H2", "util.check_login", "check_login() before request", {
            "has_token": bool(token),
            "token_length": len(token) if token else 0,
        })
        #endregion
        url = "https://qcadeltek03.qcasystems.com/vantagepoint/visionservices.asmx/GetIAccessConfiguration"
        payload = {"sessionID": token}
        headers = set_headers()
        #region agent log
        _agent_log("H2", "util.check_login", "check_login() headers from set_headers", {
            "has_token_in_headers": bool(headers.get("Token")),
            "has_cookie_in_headers": bool(headers.get("Cookie")),
        })
        #endregion
        response = requests.post(url, headers = headers, json = payload)
        #region agent log
        _agent_log("H2", "util.check_login", "check_login() after request", {
            "status_code": response.status_code,
            "response_length": len(response.text) if response.text else 0,
        })
        #endregion
        cookies = get_config(["COOKIES"])
        #region agent log
        _agent_log("H2", "util.check_login", "check_login() cookie check", {
            "has_cookies": bool(cookies),
            "has_aspnet": "ASP.NET_SessionId" in (cookies or ""),
        })
        #endregion
        if response.status_code == 200 and "ASP.NET_SessionId" in cookies:
            data = response.json()
            #region agent log
            _agent_log("H2", "util.check_login", "check_login() success", {
                "user_email": data.get("d", {}).get("UserInfo", {}).get("EMail") if isinstance(data, dict) else None,
            })
            #endregion
            #print("✅ Login Success, User: " + data["d"]["UserInfo"]["EMail"])
            #print()
            return data["d"]["UserInfo"]["EMail"]
        else: 
            #region agent log
            _agent_log("H2", "util.check_login", "check_login() failed", {
                "status_code": response.status_code,
                "has_aspnet": "ASP.NET_SessionId" in (cookies or ""),
            })
            #endregion
            return None
    except Exception as e:
        #region agent log
        _agent_log("H2", "util.check_login", "check_login() exception", {
            "error": str(e),
        })
        #endregion
        set_config("TOKEN", "")
        set_config("WWWBEARER", "")
        set_config("COOKIES", "")
        print("❌ Error in function check_login()")
        return False


def set_headers():
    from pathlib import Path
    import json
    import time
    DEBUG_LOG_PATH = Path(r"c:\cursor\.cursor\debug.log")
    def _agent_log(hypothesis_id: str, location: str, message: str, data: dict = None):
        payload = {
            "sessionId": "debug-session",
            "runId": "pre-fix",
            "hypothesisId": hypothesis_id,
            "location": location,
            "message": message,
            "data": data or {},
            "timestamp": int(time.time() * 1000),
        }
        try:
            DEBUG_LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
            with DEBUG_LOG_PATH.open("a", encoding="utf-8") as f:
                f.write(json.dumps(payload, ensure_ascii=False) + "\n")
        except Exception:
            pass
    #region agent log
    _agent_log("H3", "util.set_headers", "set_headers() entry", {})
    #endregion
    WWWBEARER = get_config(["WWWBEARER"])
    TOKEN = get_config(["TOKEN"])
    COOKIES = get_config(["COOKIES"])
    #region agent log
    _agent_log("H3", "util.set_headers", "set_headers() after get_config", {
        "has_wwwbearer": bool(WWWBEARER),
        "has_token": bool(TOKEN),
        "has_cookies": bool(COOKIES),
        "token_length": len(TOKEN) if TOKEN else 0,
        "cookie_length": len(COOKIES) if COOKIES else 0,
        "cookie_has_aspnet": "ASP.NET_SessionId" in (COOKIES or ""),
    })
    #endregion
    headers = {
        "accept": "application/json, text/javascript, */*; q=0.01",
        "Content-Type": "application/json; charset=UTF-8",
        "www-bearer": WWWBEARER,
        "Token": TOKEN,
        "Cookie": COOKIES
    }
    #region agent log
    _agent_log("H3", "util.set_headers", "set_headers() return", {
        "headers_keys": list(headers.keys()),
    })
    #endregion
    return headers

def init_workdir():
    current_path = os.getcwd()
    pattern = os.path.join(".", "job*.xlsx")

    for file_path in glob.glob(pattern):
        try:
            os.remove(file_path)
        except Exception as e:
            print(f"Failed to delete {file_path}: {e}")
    
    before, _, after = current_path.partition("OneDrive - QCA Systems Ltd")
    if not after:
        print(f'⚠️ OneDrive - QCA Systems Ltd not found in current path: {current_path}, the power automate can not been used.')
        print()
        set_config("ONEDRIVEDIR", str(current_path))
        return

    ONEDRIVEDIR = before + "OneDrive - QCA Systems Ltd"
    WORKDIR =  after.lstrip("\\")

    set_config("ONEDRIVEDIR", ONEDRIVEDIR)
    set_config("WORKDIR", WORKDIR)


def assamble_projects(projects):
    searchOptions = [{"name":"Status","value":"[IS_EMPTY]","type":"dropdown","seq":1,"tableName":"PR","opp":"!=","condition":"and","searchLevel":1,"valueDescription":""}]
    for project in projects:
        searchOptions.append({"name":"selectedResultIds","value":project,"type":"wbs1","seq":2,"searchLevel":0,"valueDescription":""})
    
    return searchOptions

def check_folder(folderName):
    if not os.path.exists(folderName):
        os.makedirs(folderName)
        print(f"📁 Folder created: {folderName}")

def clear_folder(folderName):
    onedrivedir = get_config(["ONEDRIVEDIR"])
    workdir = get_config(["WORKDIR"])

    folder_path = os.path.join(onedrivedir, workdir, folderName)
    for filename in os.listdir(folder_path):
        file_path = os.path.join(folder_path, filename)
        if os.path.isfile(file_path):
            os.remove(file_path)

def get_vendor_email(clientID):
    headers = set_headers()
    try:
        url = "https://qcadeltek03.qcasystems.com/vantagepoint/vision/Client/"+ clientID +"/Address/"
        response = requests.get(url, headers=headers  )
        data = response.json()
        email = ""
        if (len(data) == 1):
            email = data[0]["Email"]
        else:
            for d in data:
                if d["Email"] != "" and (not d["Email"] in email): 
                    if "AP Automation" in d['Address'] or "AP Mailing" in d['Address']: 
                        email = d["Email"]+";"
                        break
                    else:
                        email += d["Email"]+";"
        return email
    except Exception as e:
        print("❌ Error in function get_vendor_email()")

def get_clientID(clientName):
    try:
        url = "https://qcadeltek03.qcasystems.com/vantagepoint/visionservices.asmx/GetLookupHash"
        headers = set_headers()
        payload = {
            "sessionID": get_config(["TOKEN"]),
            "hash":{
                "filter": clientName,
                "lookuptype": "clientvendor","vendorOnly":"N"
            },
            "page": 1,
            "pagesize": 20,
            "order": "name"
        }
        response = requests.post(url, headers=headers, json=payload)
        records = response.json()
        for record in records['d']:
            if record['IsClient'] == "Y":
                return record['Key']
    except Exception as e:
        print("❌ Error in function get_clientID() with input: " + clientName)

def save_excel(wb, records):
    try:
        ws = wb.active
        table_range = f"A1:I{records}"
        if "Table1" in ws.tables:
            del ws.tables["Table1"]
        tab = Table(displayName = "Table1", ref = table_range)
        ws.add_table(tab)
        timestamp = datetime.now().strftime("%Y%m%d")
        excelName = os.path.join( f"Job_{timestamp}.xlsx")
        wb.save(excelName)

        return excelName
    except Exception as e:
        print("❌Error in function save_excel()")


def cleanup_projectID(project_id: str) -> str:
    if not project_id:
        return project_id
    replace_map = {
        "/": "[_$2F_]",
        "&": "[_$26_]",
    }
    for k, v in replace_map.items():
        project_id = project_id.replace(k, v)
    return project_id

def cleanup_projectName(projName: str) -> str:
    if not projName:
        return projName
    replace_map = {
        " ": "%20",
        "/": "%2f",
        "&": "%26",
    }
    for k, v in replace_map.items():
        projName = projName.replace(k, v)
    return projName
